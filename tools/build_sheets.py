"""Practice workbooks for Modules 2–6 (Module 1 has its own builder)."""
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, ScatterChart, Reference, Series
from sheetkit import *
from datakit import GEN

exec(GEN)
sample = make_sample_district()
messy = make_messy_record(sample)
D_RANGE = f"Data!$B$2:$B${MAXR}"; Y_RANGE = f"Data!$C$2:$C${MAXR}"; M_RANGE = f"Data!$D$2:$D${MAXR}"; DT_RANGE = f"Data!$A$2:$A${MAXR}"
YEARS = list(range(1981, 2025))


def daily_tab(wb, s, title="Data"):
    dt = wb.create_sheet(title)
    widths(dt, {"A": 13, "B": 11, "C": 8, "D": 8, "F": 70})
    hdr(dt, 1, ["Date", "Rain (mm)", "Year", "Month"])
    for i, (d, v) in enumerate(s.items()):
        r = i + 2
        dt.cell(r, 1, d.to_pydatetime()).number_format = "dd-mmm-yyyy"
        dt.cell(r, 2, None if pd.isna(v) else float(v)).number_format = "0.0"
        dt.cell(r, 3, f'=IF(A{r}="","",YEAR(A{r}))'); dt.cell(r, 4, f'=IF(A{r}="","",MONTH(A{r}))')
    text(dt, "F2", "Sample district (illustrative, synthetic), 1981–2024.", BOLD)
    text(dt, "F3", "To use your own place: paste the date and rain_mm columns from the CSV you downloaded in Session 8 into columns A and B "
                   "(starting at row 2, replacing these rows), and copy the Year and Month formulas down. Every table updates.", NOTE)
    dt.freeze_panes = "A2"
    return dt


def annual_tab(wb):
    at = wb.create_sheet("Annual")
    widths(at, {"A": 8, "B": 12, "C": 13, "D": 12, "E": 11, "F": 11, "G": 13, "H": 15, "I": 3, "J": 30, "K": 14})
    hdr(at, 1, ["Year", "Total (mm)", "Jun–Sep (mm)", "Wettest day (mm)", "Rainy days", "Heavy days", "Rain per rainy day", "Jun–Sep departure (%)"])
    for i, y in enumerate(YEARS):
        r = i + 2
        at[f"A{r}"] = y
        at[f"B{r}"] = f"=SUMIFS({D_RANGE},{Y_RANGE},A{r})"
        at[f"C{r}"] = f'=SUMIFS({D_RANGE},{Y_RANGE},A{r},{M_RANGE},">=6",{M_RANGE},"<=9")'
        at[f"D{r}"] = f"=_xlfn.MAXIFS({D_RANGE},{Y_RANGE},A{r})"
        at[f"E{r}"] = f'=COUNTIFS({Y_RANGE},A{r},{D_RANGE},">=2.5")'
        at[f"F{r}"] = f'=COUNTIFS({Y_RANGE},A{r},{D_RANGE},">=64.5")'
        at[f"G{r}"] = f"=B{r}/E{r}"
        at[f"H{r}"] = f"=(C{r}-$K$3)/$K$3*100"
        for c, fmt in zip("BCDEFGH", ["#,##0.0", "#,##0.0", "0.0", "0", "0", "0.0", "+0.0;-0.0"]):
            at[f"{c}{r}"].number_format = fmt; at[f"{c}{r}"].font = B
    last = len(YEARS) + 1
    put(at, "J2", "Normal annual (1991–2020, mm)", BOLD); put(at, "K2", f'=AVERAGEIFS(B2:B{last},A2:A{last},">=1991",A2:A{last},"<=2020")', BOLD, GRN, "#,##0.0")
    put(at, "J3", "Normal Jun–Sep (1991–2020, mm)", BOLD); put(at, "K3", f'=AVERAGEIFS(C2:C{last},A2:A{last},">=1991",A2:A{last},"<=2020")', BOLD, GRN, "#,##0.0")
    at.freeze_panes = "A2"
    return at, last


# ============================== MODULE 2 ==============================
def build_m2():
    wb = Workbook()
    start_sheet(wb, "Module 2 · Finding and Trusting Water Data — Practice workbook",
                [("Tabs", BOLD), ("S06 Datasets · S07 Three records · S08 Describe · S09 Messy record · S10 Data card", B),
                 ("Data tabs: Data (sample district daily), Three sources 2021 (illustrative), Messy record.", B),
                 ("All data here is illustrative or synthetic, for practice. Session 8 shows how to download real data for your own place.", NOTE)])
    # ---- data tabs
    daily_tab(wb, sample)
    at, last = annual_tab(wb)
    g = sample["2021"]; rng = np.random.default_rng(7)
    grid = (0.55 * g + 0.45 * g.rolling(3, center=True, min_periods=1).mean() * 0.97).round(1)
    sat = g.rolling(2, min_periods=1).mean() * 1.03
    sat[(g > 0) & (g < 3) & (rng.random(len(g)) < 0.5)] = 0.0
    sat = (sat + np.where(rng.random(len(g)) < 0.03, rng.gamma(1, 3, len(g)), 0)).round(1)
    ts = wb.create_sheet("Three sources 2021")
    widths(ts, {"A": 13, "B": 16, "C": 18, "D": 16})
    hdr(ts, 1, ["Date", "Rain gauge (point)", "Gridded (area average)", "Satellite estimate"])
    for i, d in enumerate(g.index):
        r = i + 2
        ts.cell(r, 1, d.to_pydatetime()).number_format = "dd-mmm-yyyy"
        for j, s in enumerate([g, grid, sat]): ts.cell(r, 2 + j, float(s.iloc[i])).number_format = "0.0"
    TL = len(g) + 1
    mr = wb.create_sheet("Messy record")
    widths(mr, {"A": 13, "B": 11, "C": 8, "D": 14})
    hdr(mr, 1, ["date", "rain_mm", "year", "first time?"])
    for i, row in messy.iterrows():
        r = i + 2
        mr.cell(r, 1, row.date); mr.cell(r, 2, None if pd.isna(row.rain_mm) else float(row.rain_mm))
        mr.cell(r, 3, f"=VALUE(LEFT(A{r},4))"); mr.cell(r, 4, f"=COUNTIF($A$2:A{r},A{r})=1")
    ML = len(messy) + 1

    # ---- S06
    ws = wb.create_sheet("S06 Datasets")
    widths(ws, {"A": 3, "B": 5, "C": 70, "D": 30, "E": 14})
    text(ws, "B2", "Session 6 · Which dataset first?", TITLE)
    DS = ["IMD district rainfall", "State telemetric network", "CWC flood forecasts", "CWC reservoir bulletin",
          "IMD gridded rainfall", "NASA POWER", "NASA GPM IMERG", "CGWB observation wells"]
    items = [("Q1", "A newspaper wants to know whether your district's monsoon was deficient this year.", "txt", "IMD district rainfall", DS, None, None),
             ("Q2", "A panchayat wants yesterday's rainfall for its own village.", "txt", "State telemetric network", DS, None, None),
             ("Q3", "A school by the river wants to know if the river will cross the danger level tomorrow.", "txt", "CWC flood forecasts", DS, None, None),
             ("Q4", "A farmer group wants to know whether the upstream reservoir is fuller than last year, this week.", "txt", "CWC reservoir bulletin", DS, None, None),
             ("Q5", "A historian wants to compare rainfall in the 1918 drought with today.", "txt", "IMD gridded rainfall", DS, None, None),
             ("Q6", "A researcher needs daily rainfall since 1981 for a hill village with no gauge, in minutes.", "txt", "NASA POWER", DS, None, None),
             ("Q7", "A borewell owner wants to know if groundwater in the district has been falling.", "txt", "CGWB observation wells", DS, None, None),
             ("Q8", "IMD's day ends at 8:30 am IST; a global dataset's day ends at midnight UTC. Which metadata item warns you?", "txt", "Day definition",
              ["Unit", "Day definition", "Licence", "Missing-value flag"], None, None)]
    questions(ws, 4, items)

    # ---- S07
    ws = wb.create_sheet("S07 Three records")
    widths(ws, {"A": 3, "B": 5, "C": 62, "D": 22, "E": 14, "F": 3, "G": 24, "H": 13, "I": 13, "J": 13})
    text(ws, "B2", "Session 7 · Stations, grids and satellites (illustrative 2021 data)", TITLE)
    hdr(ws, 4, ["Record", "Year total (mm)", "Wettest day (mm)", "Rainy days"], col=7)
    for i, (name, col) in enumerate([("Rain gauge (point)", "B"), ("Gridded (area average)", "C"), ("Satellite estimate", "D")]):
        r = 5 + i; rng_ = f"'Three sources 2021'!{col}2:{col}{TL}"
        put(ws, f"G{r}", name); put(ws, f"H{r}", f"=SUM({rng_})", BOLD, GRN, "#,##0.0")
        put(ws, f"I{r}", f"=MAX({rng_})", BOLD, GRN, "0.0"); put(ws, f"J{r}", f'=COUNTIF({rng_},">=2.5")', BOLD, GRN, "0")
    peak = 'INDEX($G$5:$G$7,MATCH(MAX($I$5:$I$7),$I$5:$I$7,0))'
    items = [("Q1", "2021 year total for the rain gauge (mm)?", "num", "H5", None, "#,##0.0", None),
             ("Q2", "% difference of satellite total from gauge total: (satellite − gauge) ÷ gauge × 100.", "num", "(H7-H5)/H5*100", None, "0.0", None),
             ("Q3", "Which record has the highest wettest day?", "txt", "=" + peak, ["Rain gauge (point)", "Gridded (area average)", "Satellite estimate"], None, None),
             ("Q4", "How many mm lower is the gridded wettest day than the gauge wettest day?", "num", "I5-I6", None, "0.0", None),
             ("Q5", "Why does a grid usually show a smaller wettest day than one gauge?", "txt", "Grid values average a large area",
              ["Grids use older instruments", "Grid values average a large area", "Grids record every other day"], None, None),
             ("Q6", "Best source for rain over the Bay of Bengal during a cyclone?", "txt", "Satellite estimate", ["Rain gauge", "IMD gridded rainfall", "Satellite estimate"], None, None)]
    questions(ws, 9, items)

    # ---- S08
    ws = wb.create_sheet("S08 Describe")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 8 · Describe the record (Data tab)", TITLE)
    items = [("Q1", "How many days are in the record?", "num", f"COUNT({DT_RANGE})", None, "#,##0", None),
             ("Q2", "In which year does the record start?", "num", f"MIN({Y_RANGE})", None, "0", None),
             ("Q3", f"Average annual total over all years (mm)? (Use the Annual tab.)", "num", f"AVERAGE(Annual!B2:B{last})", None, "#,##0.0", None),
             ("Q4", "Which year had the highest annual total?", "num", f"INDEX(Annual!A2:A{last},MATCH(MAX(Annual!B2:B{last}),Annual!B2:B{last},0))", None, "0", None),
             ("Q5", "What % of days had exactly 0 mm?", "num", f'COUNTIF({D_RANGE},0)/COUNT({D_RANGE})*100', None, "0.0", None),
             ("Q6", "Most important item on a data card for NASA POWER data?", "txt", "The source and its spatial detail",
              ["The file size", "The source and its spatial detail", "The chart colour"], None, None)]
    questions(ws, 4, items)

    # ---- S09
    ws = wb.create_sheet("S09 Messy record")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 9 · Investigate the Messy record tab", TITLE)
    R = f"'Messy record'!$B$2:$B${ML}"; A_ = f"'Messy record'!$A$2:$A${ML}"; Y = f"'Messy record'!$C$2:$C${ML}"; FT = f"'Messy record'!$D$2:$D${ML}"
    items = [("Q1", "How many rows contain −999?", "num", f"COUNTIF({R},-999)", None, "0", None),
             ("Q2", "How many blank cells in rain_mm?", "num", f"COUNTBLANK({R})", None, "0", None),
             ("Q3", "How many impossible values (below 0 or above 500, not counting −999)?", "num", f'COUNTIFS({R},"<0",{R},"<>-999")+COUNTIF({R},">500")', None, "0", None),
             ("Q4", "How many extra duplicate rows (dates appearing a second time)?", "num", f'COUNTIF({FT},FALSE)', None, "0", None),
             ("Q5", "How many days are missing entirely (no row at all)? Hint: last date − first date + 1 − number of distinct dates.", "num",
              f"(DATEVALUE('Messy record'!A{ML})-DATEVALUE('Messy record'!A2)+1)-COUNTIF({FT},TRUE)", None, "0", None),
             ("Q6", "Clean 2019 total: values from 0 to 500 only, first occurrence of each date only (mm)?", "num",
              f'SUMIFS({R},{Y},2019,{R},">=0",{R},"<=500",{FT},TRUE)', None, "#,##0.0", None),
             ("Q7", "186.2 mm on one day, with heavy rain the days before. What should you do?", "txt", "Keep it after checking neighbours and another source",
              ["Delete it as an outlier", "Keep it after checking neighbours and another source", "Set it to 0"], None, None)]
    questions(ws, 4, items)

    # ---- S10 data card template
    ws = wb.create_sheet("S10 Data card")
    widths(ws, {"A": 3, "B": 40, "C": 60})
    text(ws, "B2", "Session 10 · Workshop: my district's data card (capstone step 2)", TITLE)
    text(ws, "B3", "Fill this in from the Session 10 notebook (or paste your district into the Data tab and use the Annual tab).", NOTE)
    hdr(ws, 5, ["Item", "My district"], col=2)
    for i, item in enumerate(["District and state", "Source (dataset and parameter)", "Location used (lat, lon)", "Period", "Spatial detail",
                              "Missing days / years below 95% complete", "1991–2020 average (mm)", "IMD normal (mm) and difference (%)",
                              "Wettest day (mm, date)", "Problems found", "My decision about using this data"]):
        r = 6 + i; put(ws, f"B{r}", item); put(ws, f"C{r}", None, INP, YEL, wrap=True); ws.row_dimensions[r].height = 30
    wb.move_sheet("Data", offset=len(wb.sheetnames)); wb.move_sheet("Annual", offset=len(wb.sheetnames))
    wb.move_sheet("Three sources 2021", offset=len(wb.sheetnames)); wb.move_sheet("Messy record", offset=len(wb.sheetnames))
    finish(wb, "sheets/WDL_M2_practice.xlsx")


# ============================== MODULE 3 ==============================
def build_m3():
    wb = Workbook()
    start_sheet(wb, "Module 3 · Reading Rainfall — Practice workbook",
                [("Tabs", BOLD), ("S11 Normal · S12 Anomalies · S13 Extremes · S14 Return periods · S15 Profile", B),
                 ("The Data tab holds the sample district (synthetic, 1981–2024). Paste your own district's downloaded CSV there to analyse it instead.", NOTE)])
    daily_tab(wb, sample)
    at, last = annual_tab(wb)
    # IMD category next to departure
    at["I1"] = None
    hdr(at, 1, ["Category"], col=13)
    widths(at, {"M": 16})
    for r in range(2, last + 1):
        at[f"M{r}"] = f'=IF(H{r}<=-60,"Large deficient",IF(H{r}<=-20,"Deficient",IF(H{r}<20,"Normal",IF(H{r}<60,"Excess","Large excess"))))'
    B_ = f"Annual!$B$2:$B${last}"; C_ = f"Annual!$C$2:$C${last}"; D_ = f"Annual!$D$2:$D${last}"; A_ = f"Annual!$A$2:$A${last}"; H_ = f"Annual!$H$2:$H${last}"
    CATS = ["Large excess", "Excess", "Normal", "Deficient", "Large deficient"]

    ws = wb.create_sheet("S11 Normal")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 11 · Totals, averages and the normal", TITLE)
    items = [("Q1", "1991–2020 normal annual rainfall (mm)?", "num", "Annual!$K$2", None, "#,##0.0", None),
             ("Q2", "Median annual total, all years (mm)?", "num", f"MEDIAN({B_})", None, "#,##0.0", None),
             ("Q3", "Which year was the driest?", "num", f"INDEX({A_},MATCH(MIN({B_}),{B_},0))", None, "0", None),
             ("Q4", "Average % of the year's rain falling June–September?", "num", f"AVERAGE({C_})/AVERAGE({B_})*100", None, "0.0", None),
             ("Q5", "Is the mean higher or lower than the median?", "txt", f'=IF(AVERAGE({B_})>MEDIAN({B_}),"Higher","Lower")', ["Higher", "Lower"], None, None)]
    nxt = questions(ws, 4, items)
    ch = BarChart(); ch.type = "col"; ch.title = "Annual rainfall (mm)"; ch.legend = None
    ch.add_data(Reference(at, min_col=2, min_row=1, max_row=last), titles_from_data=True); ch.set_categories(Reference(at, min_col=1, min_row=2, max_row=last))
    ch.height, ch.width = 8, 22; ws.add_chart(ch, f"B{nxt}")

    ws = wb.create_sheet("S12 Anomalies")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 12 · Anomalies (June–September)", TITLE)
    r02 = YEARS.index(2002) + 2
    items = [("Q1", "June–September normal, 1991–2020 (mm)?", "num", "Annual!$K$3", None, "#,##0.0", None),
             ("Q2", "Departure (%) of the 2002 monsoon from normal?", "num", f"Annual!H{r02}", None, "0.0", None),
             ("Q3", "IMD category of the 2002 monsoon?", "txt", f"=Annual!M{r02}", CATS, None, None),
             ("Q4", "How many monsoons were Deficient or Large deficient (−20% or lower)?", "num", f'COUNTIF({H_},"<=-20")', None, "0", None),
             ("Q5", "Which year had the largest positive departure?", "num", f"INDEX({A_},MATCH(MAX({H_}),{H_},0))", None, "0", None),
             ("Q6", '"Monsoon 15% below normal": IMD category?', "txt", "Normal", CATS, None, None)]
    nxt = questions(ws, 4, items)
    ch = BarChart(); ch.type = "col"; ch.title = "June–September departure from normal (%)"; ch.legend = None
    ch.add_data(Reference(at, min_col=8, min_row=1, max_row=last), titles_from_data=True); ch.set_categories(Reference(at, min_col=1, min_row=2, max_row=last))
    ch.height, ch.width = 8, 22; ws.add_chart(ch, f"B{nxt}")

    ws = wb.create_sheet("S13 Extremes")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 13 · Extremes and intensity (Data tab)", TITLE)
    items = [("Q1", "Days with 64.5 mm or more?", "num", f'COUNTIF({D_RANGE},">=64.5")', None, "0", None),
             ("Q2", "Days with 115.6 mm or more?", "num", f'COUNTIF({D_RANGE},">=115.6")', None, "0", None),
             ("Q3", "% of all rain on days of 64.5 mm or more?", "num", f'SUMIF({D_RANGE},">=64.5")/SUM({D_RANGE})*100', None, "0.0", None),
             ("Q4", "Wettest day in the record (mm)?", "num", f"MAX({D_RANGE})", None, "0.0", None),
             ("Q5", "72 mm in 24 hours: IMD term?", "txt", "Heavy rain", ["Moderate rain", "Heavy rain", "Very heavy rain", "Extremely heavy rain"], None, None),
             ("Q6", "% of rainy days (above 0) that were very light (0.1–2.4 mm)?", "num", f'COUNTIFS({D_RANGE},">0",{D_RANGE},"<2.5")/COUNTIF({D_RANGE},">0")*100', None, "0.0", None)]
    questions(ws, 4, items)

    ws = wb.create_sheet("S14 Return periods")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14, "F": 3, "G": 8, "H": 16, "I": 16})
    text(ws, "B2", "Session 14 · Return periods (annual maxima from the Annual tab)", TITLE)
    hdr(ws, 4, ["Rank", "Wettest day (mm)", "T = (n+1)/rank"], col=7)
    n = len(YEARS)
    for k in range(1, n + 1):
        r = 4 + k
        put(ws, f"G{r}", k); put(ws, f"H{r}", f"=LARGE({D_},G{r})", fmt="0.0"); put(ws, f"I{r}", f"=(COUNT({D_})+1)/G{r}", fmt="0.0")
    items = [("Q1", "Return period of the largest annual maximum (years)?", "num", f"COUNT({D_})+1", None, "0.0", None),
             ("Q2", "Return period of the rank 5 event (years)?", "num", f"(COUNT({D_})+1)/5", None, "0.0", None),
             ("Q3", "Rain on the rank 5 wettest day (mm)?", "num", f"LARGE({D_},5)", None, "0.0", None),
             ("Q4", "Chance (%) of a 50-year event in any single year?", "num", "100/50", None, "0.0", None),
             ("Q5", "Chance (%) of at least one 10-year event in the next 10 years?", "num", "(1-(1-1/10)^10)*100", None, "0.0", None),
             ("Q6", "A 100-year storm hit last year. Are we safe for 99 years?", "txt", "No: the chance is still 1% each year",
              ["Yes: it cannot happen again soon", "No: the chance is still 1% each year"], None, None)]
    questions(ws, 4, items, col_letters=("B", "C", "D", "E"))

    ws = wb.create_sheet("S15 Profile")
    widths(ws, {"A": 3, "B": 46, "C": 22, "D": 50})
    text(ws, "B2", "Session 15 · Workshop: rainfall profile (capstone step 3)", TITLE)
    text(ws, "B3", "Paste your own district into the Data tab first. Everything below updates.", NOTE)
    hdr(ws, 5, ["Fact", "Value", "Notes"], col=2)
    facts = [("Normal annual rainfall, 1991–2020 (mm)", "=Annual!K2", "#,##0"), ("June–September normal (mm)", "=Annual!K3", "#,##0"),
             ("Monsoon share of the year (%)", f"=AVERAGE({C_})/AVERAGE({B_})*100", "0"),
             ("Wettest year", f"=INDEX({A_},MATCH(MAX({B_}),{B_},0))", "0"), ("Driest year", f"=INDEX({A_},MATCH(MIN({B_}),{B_},0))", "0"),
             ("Monsoons: excess or more", f'=COUNTIF({H_},">=20")', "0"), ("Monsoons: normal", f'=COUNTIFS({H_},">-20",{H_},"<20")', "0"),
             ("Monsoons: deficient or less", f'=COUNTIF({H_},"<=-20")', "0"),
             ("Heavy-rain days per decade", f'=COUNTIF({D_RANGE},">=64.5")/(COUNT({B_})/10)', "0.0"),
             ("Wettest day on record (mm)", f"=MAX({D_RANGE})", "0.0")]
    for i, (k, f, fmt) in enumerate(facts):
        r = 6 + i; put(ws, f"B{r}", k); put(ws, f"C{r}", f, BOLD, GRN, fmt)
    put(ws, "B17", "My one-sentence summary"); ws.merge_cells("C17:D17"); put(ws, "C17", None, INP, YEL, wrap=True); ws.row_dimensions[17].height = 40
    for s_ in ["Data", "Annual"]: wb.move_sheet(s_, offset=len(wb.sheetnames))
    finish(wb, "sheets/WDL_M3_practice.xlsx")


# ============================== MODULE 4 ==============================
def build_m4():
    wb = Workbook()
    start_sheet(wb, "Module 4 · Rivers, Reservoirs and Floods — Practice workbook",
                [("Tabs", BOLD), ("S16 Hydrograph · S17 Rating curve · S18 Reservoir · S19 Rational method · S20 Event card", B),
                 ("The hydrograph and reservoir are illustrative teaching models, the same as in the notebooks.", NOTE)])
    # hydrograph (same model as the notebook)
    PATTERN = np.array([0.10, 0.20, 0.30, 0.20, 0.12, 0.08]); hours = 96; t = np.arange(hours)
    rain_h = np.zeros(hours); rain_h[2:8] = 80 * PATTERN; eff = rain_h * 0.4
    tp = max(2.0, 1.0 + 0.35 * np.sqrt(250)); tb = 2.67 * tp; qp = 2 * 250 * 1000 / (tb * 3600)
    uh = np.interp(t, [0, tp, tb], [0, qp, 0], right=0); q = np.convolve(eff, uh)[:hours] + 5.0
    ws = wb.create_sheet("S16 Hydrograph")
    widths(ws, {"A": 3, "B": 5, "C": 62, "D": 16, "E": 14, "F": 3, "G": 8, "H": 11, "I": 12})
    text(ws, "B2", "Session 16 · Read the hydrograph: 80 mm over 250 km², runoff coefficient 0.4", TITLE)
    hdr(ws, 4, ["Hour", "Rain (mm)", "Flow (m³/s)"], col=7)
    for i in range(hours):
        r = 5 + i; ws[f"G{r}"] = i; ws[f"H{r}"] = float(round(rain_h[i], 2)); ws[f"I{r}"] = float(round(q[i], 2))
    HL = 4 + hours
    items = [("Q1", "Peak flow (m³/s)?", "num", f"MAX(I5:I{HL})", None, "#,##0", None),
             ("Q2", "Hour of the peak?", "num", f"INDEX(G5:G{HL},MATCH(MAX(I5:I{HL}),I5:I{HL},0))", None, "0", None),
             ("Q3", "Runoff volume (MCM) = coefficient × rain (m) × area (m²) ÷ 1,000,000?", "num", "0.4*0.08*250000000/1000000", None, "0.00", None),
             ("Q4", "If the runoff coefficient doubles to 0.8, the peak will…", "txt", "Roughly double", ["Stay the same", "Roughly double", "Halve"], None, None),
             ("Q5", "A 1,000 km² catchment vs 50 km², same storm: lag time is…", "txt", "Longer", ["Shorter", "Longer", "The same"], None, None)]
    questions(ws, 4, items)
    lc = LineChart(); lc.title = "Hydrograph"; lc.legend = None; lc.y_axis.title = "m³/s"; lc.x_axis.title = "Hour"
    lc.add_data(Reference(ws, min_col=9, min_row=4, max_row=HL), titles_from_data=True); lc.set_categories(Reference(ws, min_col=7, min_row=5, max_row=HL))
    lc.height, lc.width = 7.5, 16; ws.add_chart(lc, "B13")

    ws = wb.create_sheet("S17 Rating curve")
    widths(ws, {"A": 3, "B": 5, "C": 62, "D": 16, "E": 14, "F": 3, "G": 22, "H": 12})
    text(ws, "B2", "Session 17 · Rating curve Q = a × (h − h0)^b", TITLE)
    put(ws, "G4", "a"); put(ws, "H4", 40); put(ws, "G5", "h0 (m)"); put(ws, "H5", 1.2); put(ws, "G6", "b"); put(ws, "H6", 1.6)
    put(ws, "G8", "Calculator: level (m)", BOLD); put(ws, "H8", 4, INP, YEL); put(ws, "G9", "Flow (m³/s)", BOLD); put(ws, "H9", "=$H$4*MAX(H8-$H$5,0)^$H$6", BOLD, GRN, "#,##0")
    items = [("Q1", "60 m wide × 2.5 m deep × 1.2 m/s: flow (m³/s)?", "num", "60*2.5*1.2", None, "#,##0", None),
             ("Q2", "Flow at a gauge level of 3.0 m (m³/s)?", "num", "$H$4*(3-$H$5)^$H$6", None, "#,##0", None),
             ("Q3", "Flow at the danger level, 6.0 m (m³/s)?", "num", "$H$4*(6-$H$5)^$H$6", None, "#,##0", None),
             ("Q4", "Level 2 m → 4 m: how many times larger is the flow?", "num", "((4-$H$5)/(2-$H$5))^$H$6", None, "0.0", None),
             ("Q5", "Level for 1,000 m³/s (m)? h = h0 + (Q ÷ a)^(1/b)", "num", "$H$5+(1000/$H$4)^(1/$H$6)", None, "0.00", None),
             ("Q6", "Level for a 50,000 cusec release (m)? (÷ 35.3147 first)", "num", "$H$5+((50000/35.3147)/$H$4)^(1/$H$6)", None, "0.00", None)]
    questions(ws, 11, items)

    # reservoir (same model as notebook)
    s = sample.fillna(0)["2013-06-01":"2024-12-31"]
    inflow = s / 1000 * 3000 * 1e6 * 0.31 / 1e6
    live, rows = 300.0, []
    for d, qq in inflow.items():
        live = live + qq - (1.0 if 6 <= d.month <= 10 else 3.6)
        sp = max(live - 1000, 0); live = min(max(live, 0), 1000); rows.append((d, qq, live, sp))
    rs = wb.create_sheet("Reservoir data")
    widths(rs, {"A": 13, "B": 12, "C": 14, "D": 12, "E": 12})
    hdr(rs, 1, ["Date", "Inflow (MCM)", "Live storage (MCM)", "Spill (MCM)", "% of live capacity"])
    for i, (d, qq, lv, sp) in enumerate(rows):
        r = i + 2
        rs.cell(r, 1, d.to_pydatetime()).number_format = "dd-mmm-yyyy"; rs.cell(r, 2, round(qq, 3)); rs.cell(r, 3, round(lv, 3)); rs.cell(r, 4, round(sp, 3))
        rs.cell(r, 5, f"=C{r}/1000*100").number_format = "0.0"
    RL = len(rows) + 1
    ws = wb.create_sheet("S18 Reservoir")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 16, "E": 14})
    text(ws, "B2", "Session 18 · Illustrative reservoir: gross 1,200 MCM, dead 200 MCM", TITLE)
    look = lambda d, col: f"INDEX('Reservoir data'!{col}2:{col}{RL},MATCH(DATE({d}),'Reservoir data'!A2:A{RL},0))"
    items = [("Q1", "Live capacity (MCM)?", "num", "1200-200", None, "#,##0", None),
             ("Q2", "Live storage on 15 Jan 2024 (MCM)?", "num", look("2024,1,15", "C"), None, "#,##0.0", None),
             ("Q3", "% of live capacity on 15 Jan 2024?", "num", look("2024,1,15", "E"), None, "0.0", None),
             ("Q4", "Percentage points above (+) or below (−) 15 Jan 2023?", "num", f"{look('2024,1,15', 'E')}-{look('2023,1,15', 'E')}", None, "0.0", None),
             ("Q5", "Total inflow 1 Jun – 30 Sep 2023 (MCM)?", "num", f"SUMIFS('Reservoir data'!B2:B{RL},'Reservoir data'!A2:A{RL},\">=\"&DATE(2023,6,1),'Reservoir data'!A2:A{RL},\"<=\"&DATE(2023,9,30))", None, "#,##0.0", None),
             ("Q6", "Live capacity (1,000 MCM) in TMC?", "num", "1000/28.316846592", None, "0.0", None)]
    nxt = questions(ws, 4, items)
    lc = LineChart(); lc.title = "% of live capacity"; lc.legend = None
    lc.add_data(Reference(rs, min_col=5, min_row=1, max_row=RL), titles_from_data=True); lc.set_categories(Reference(rs, min_col=1, min_row=2, max_row=RL))
    lc.x_axis.number_format = "yyyy"; lc.height, lc.width = 7.5, 20; ws.add_chart(lc, f"B{nxt}")

    ws = wb.create_sheet("S19 Rational method")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 16, "E": 14, "F": 3, "G": 26, "H": 12})
    text(ws, "B2", "Session 19 · Q = C × i × A ÷ 3.6", TITLE)
    for i, (k, v) in enumerate([("C (runoff coefficient)", 0.6), ("i (mm per hour)", 40), ("A (km²)", 2), ("Drain capacity (m³/s)", 10)]):
        put(ws, f"G{4+i}", k); put(ws, f"H{4+i}", v, INP, YEL)
    put(ws, "G8", "Peak runoff (m³/s)", BOLD); put(ws, "H8", "=H4*H5*H6/3.6", BOLD, GRN, "0.0")
    put(ws, "G9", "Does it flood?", BOLD); put(ws, "H9", '=IF(H8>H7,"FLOODS","Drain copes")', BOLD, GRN)
    items = [("Q1", "C = 0.3, i = 40 mm/h, A = 2 km²: peak (m³/s)?", "num", "0.3*40*2/3.6", None, "0.00", None),
             ("Q2", "Same, C = 0.85: peak (m³/s)?", "num", "0.85*40*2/3.6", None, "0.00", None),
             ("Q3", "60% paved (0.85) + 40% green (0.20): combined C?", "num", "0.6*0.85+0.4*0.2", None, "0.00", None),
             ("Q4", "120 mm in 3 hours: intensity (mm/h)?", "num", "120/3", None, "0", None),
             ("Q5", "Drain 10 m³/s; A = 2 km², C = 0.85, 40 mm/h. Overflow?", "txt", "Yes", ["Yes", "No"], None, None),
             ("Q6", "Four days of rain over 20,000 km², river rises over two days: type?", "txt", "River flood", ["Flash flood", "River flood", "Urban flood"], None, None)]
    questions(ws, 11, items)

    ws = wb.create_sheet("S20 Event card")
    widths(ws, {"A": 3, "B": 46, "C": 22, "D": 40})
    text(ws, "B2", "Session 20 · Workshop: extreme event card (capstone step 4)", TITLE)
    hdr(ws, 4, ["Input", "Value", "Notes"], col=2)
    inputs = [("Event date", None, "From the Session 20 notebook"), ("Rain on the day (mm)", 150, ""), ("5-day total (mm)", 210, ""),
              ("Normal for that month (mm)", 180, ""), ("District area (km²)", 4000, ""), ("Share of the day's rain in the worst hour", 0.35, "Assumption"),
              ("Town area (km²)", 5, ""), ("Town runoff coefficient", 0.7, "")]
    for i, (k, v, nte) in enumerate(inputs):
        r = 5 + i; put(ws, f"B{r}", k); put(ws, f"C{r}", v, INP, YEL); put(ws, f"D{r}", nte, NOTE)
    hdr(ws, 14, ["Result", "Value", ""], col=2)
    res = [("IMD term", '=IF(C6>=204.5,"Extremely heavy",IF(C6>=115.6,"Very heavy",IF(C6>=64.5,"Heavy","Moderate or less")))', None),
           ("% of the month's normal in one day", "=C6/C8*100", "0"), ("Water on the district that day (MCM)", "=C6/1000*C9", "#,##0"),
           ("Worst-hour intensity (mm/h)", "=C6*C10", "0"), ("Town peak runoff (m³/s)", "=C12*C6*C10*C11/3.6", "#,##0")]
    for i, (k, f, fmt) in enumerate(res):
        r = 15 + i; put(ws, f"B{r}", k); put(ws, f"C{r}", f, BOLD, GRN, fmt)
    put(ws, "B21", "What happened (news or local knowledge)"); ws.merge_cells("C21:D21"); put(ws, "C21", None, INP, YEL, wrap=True); ws.row_dimensions[21].height = 45
    wb.move_sheet("Reservoir data", offset=len(wb.sheetnames))
    finish(wb, "sheets/WDL_M4_practice.xlsx")


# ============================== MODULE 5 ==============================
def build_m5():
    wb = Workbook()
    start_sheet(wb, "Module 5 · Climate Signals in Water Data — Practice workbook",
                [("Tabs", BOLD), ("S21 Variability · S22 Trends · S23 Extremes and dry spells · S24 Claims · S25 Trend card", B),
                 ("Spreadsheets use the least-squares slope (SLOPE). The notebooks also give Sen's slope and the Mann–Kendall p-value, "
                  "which spreadsheets cannot compute easily. Use the notebook for your final verdict.", NOTE)])
    daily_tab(wb, sample)
    at, last = annual_tab(wb)
    # moving average and abs change
    hdr(at, 1, ["11-yr moving avg", "Change from last year"], col=13); widths(at, {"M": 15, "N": 16})
    for r in range(2, last + 1):
        if r - 5 >= 2 and r + 5 <= last: at[f"M{r}"] = f"=AVERAGE(B{r-5}:B{r+5})"
        if r > 2: at[f"N{r}"] = f"=ABS(B{r}-B{r-1})"
    B_ = f"Annual!$B$2:$B${last}"; A_ = f"Annual!$A$2:$A${last}"; E_ = f"Annual!$E$2:$E${last}"; F_ = f"Annual!$F$2:$F${last}"; G_ = f"Annual!$G$2:$G${last}"
    yr = lambda y: YEARS.index(y) + 2

    ws = wb.create_sheet("S21 Variability")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 21 · Trend vs variability", TITLE)
    items = [("Q1", "Standard deviation of annual totals (mm)?", "num", f"STDEV({B_})", None, "#,##0.0", None),
             ("Q2", "Coefficient of variation (SD ÷ mean × 100, %)?", "num", f"STDEV({B_})/AVERAGE({B_})*100", None, "0.0", None),
             ("Q3", "Years outside mean ± 1 SD?", "num", f'COUNTIF({B_},"<"&(AVERAGE({B_})-STDEV({B_})))+COUNTIF({B_},">"&(AVERAGE({B_})+STDEV({B_})))', None, "0", None),
             ("Q4", "Year with the biggest change from the year before?", "num", f"INDEX(Annual!$A$3:$A${last},MATCH(MAX(Annual!$N$3:$N${last}),Annual!$N$3:$N${last},0))", None, "0", None),
             ("Q5", "Average rainy days 2015–2024 minus 1981–1990?", "num", f"AVERAGE(Annual!E{yr(2015)}:E{yr(2024)})-AVERAGE(Annual!E{yr(1981)}:E{yr(1990)})", None, "0.0", None),
             ("Q6", "The wettest year in a decade proves…", "txt", "Very little: one year is natural variability", ["Rainfall is increasing", "Very little: one year is natural variability", "The climate is stable"], None, None)]
    nxt = questions(ws, 4, items)
    lc = LineChart(); lc.title = "Annual total and 11-year moving average (mm)"
    lc.add_data(Reference(at, min_col=2, min_row=1, max_row=last), titles_from_data=True)
    lc.add_data(Reference(at, min_col=13, min_row=1, max_row=last), titles_from_data=True)
    lc.set_categories(Reference(at, min_col=1, min_row=2, max_row=last)); lc.height, lc.width = 8, 22; ws.add_chart(lc, f"B{nxt}")

    ws = wb.create_sheet("S22 Trends")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 22 · How many years are enough? (least-squares slope × 10 = per decade)", TITLE)
    items = [("Q1", "Slope of the annual total, 1981–2024 (mm per decade)?", "num", f"SLOPE({B_},{A_})*10", None, "0.0", None),
             ("Q2", "Slope of the annual total, 2011–2020 (mm per decade)?", "num", f"SLOPE(Annual!B{yr(2011)}:B{yr(2020)},Annual!A{yr(2011)}:A{yr(2020)})*10", None, "0.0", None),
             ("Q3", "Slope of rainy days, 1981–2024 (days per decade)?", "num", f"SLOPE({E_},{A_})*10", None, "0.00", None),
             ("Q4", "Slope of rain per rainy day, 1981–2024 (mm per decade)?", "num", f"SLOPE({G_},{A_})*10", None, "0.00", None),
             ("Q5", "Why use at least 30 years?", "txt", "Short periods are dominated by natural swings", ["Older data is more accurate", "Short periods are dominated by natural swings", "Computers need 30 values"], None, None)]
    questions(ws, 4, items)

    ws = wb.create_sheet("S23 Extremes")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 18, "E": 14})
    text(ws, "B2", "Session 23 · Extremes and dry spells (Annual tab)", TITLE)
    items = [("Q1", "Average rain per rainy day, 2015–2024 (mm)?", "num", f"AVERAGE(Annual!G{yr(2015)}:G{yr(2024)})", None, "0.0", None),
             ("Q2", "Change in rain per rainy day, 1981–1990 to 2015–2024 (mm)?", "num", f"AVERAGE(Annual!G{yr(2015)}:G{yr(2024)})-AVERAGE(Annual!G{yr(1981)}:G{yr(1990)})", None, "0.0", None),
             ("Q3", "Heavy-rain days (≥ 64.5 mm) in 2015–2024 in total?", "num", f"SUM(Annual!F{yr(2015)}:F{yr(2024)})", None, "0", None),
             ("Q4", "Heavy-rain days in 1981–1990 in total?", "num", f"SUM(Annual!F{yr(1981)}:F{yr(1990)})", None, "0", None),
             ("Q5", "Normal rain, but reservoirs and wells low after dry years: which drought?", "txt", "Hydrological drought", ["Meteorological drought", "Agricultural drought", "Hydrological drought"], None, None)]
    questions(ws, 4, items)

    ws = wb.create_sheet("S24 Claims")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 24, "E": 14, "F": 3, "G": 12, "H": 18})
    text(ws, "B2", "Session 24 · Reading climate claims critically", TITLE)
    hdr(ws, 4, ["First year", "10-yr slope (mm/decade)"], col=7)
    for i, y in enumerate(range(1981, 2016)):
        r = 5 + i; put(ws, f"G{r}", y); put(ws, f"H{r}", f"=SLOPE(Annual!B{yr(y)}:B{yr(y)+9},Annual!A{yr(y)}:A{yr(y)+9})*10", fmt="0.0")
    P = ["Cherry-picked period", "Truncated axis", "Two axes implying a link", "Wrong unit"]
    items = [("Q1", "Bars of five years with a y-axis from 950 to 1,050 mm.", "txt", "Truncated axis", P, None, None),
             ("Q2", '2012–2018 only, captioned "rainfall has been falling for years".', "txt", "Cherry-picked period", P, None, None),
             ("Q3", "Rainfall and phone users on two y-axes, scaled to match.", "txt", "Two axes implying a link", P, None, None),
             ("Q4", "Most negative 10-year slope (mm per decade)?", "num", "MIN(H5:H39)", None, "0", None),
             ("Q5", "Full-record slope (mm per decade)?", "num", f"SLOPE({B_},{A_})*10", None, "0.0", None),
             ("Q6", 'IPCC "likely" means…', "txt", "66–100%", ["33–66%", "66–100%", "90–100%"], None, None)]
    questions(ws, 4, items, col_letters=("B", "C", "D", "E"))

    ws = wb.create_sheet("S25 Trend card")
    widths(ws, {"A": 3, "B": 40, "C": 22, "D": 50})
    text(ws, "B2", "Session 25 · Workshop: trend card (capstone step 5)", TITLE)
    text(ws, "B3", "Paste your district into the Data tab. Copy the Sen's slopes and p-values from the notebook for your final verdict.", NOTE)
    hdr(ws, 5, ["Indicator", "Least-squares slope per decade", "Verdict from the notebook (p-value)"], col=2)
    for i, (k, rng_) in enumerate([("Annual total (mm)", B_), ("Rainy days", E_), ("Rain per rainy day (mm)", G_), ("Heavy-rain days", F_)]):
        r = 6 + i; put(ws, f"B{r}", k); put(ws, f"C{r}", f"=SLOPE({rng_},{A_})*10", BOLD, GRN, "+0.00;-0.00"); put(ws, f"D{r}", None, INP, YEL)
    put(ws, "B11", "My verdict in words"); ws.merge_cells("C11:D11"); put(ws, "C11", None, INP, YEL, wrap=True); ws.row_dimensions[11].height = 45
    put(ws, "B12", "Caveats"); ws.merge_cells("C12:D12"); put(ws, "C12", None, INP, YEL, wrap=True); ws.row_dimensions[12].height = 45
    for s_ in ["Data", "Annual"]: wb.move_sheet(s_, offset=len(wb.sheetnames))
    finish(wb, "sheets/WDL_M5_practice.xlsx")


# ============================== MODULE 6 ==============================
def build_m6():
    wb = Workbook()
    start_sheet(wb, "Module 6 · From Data to Dashboard — Practice workbook",
                [("Tabs", BOLD), ("S26 Honest charts · S27 Chart choice · S28 Dashboard · S29 Story checker", B),
                 ("The Dashboard tab updates when you paste your own district into the Data tab.", NOTE)])
    daily_tab(wb, sample)
    at, last = annual_tab(wb)
    B_ = f"Annual!$B$2:$B${last}"; H_ = f"Annual!$H$2:$H${last}"; E_ = f"Annual!$E$2:$E${last}"
    ws = wb.create_sheet("S26 Honest charts")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 34, "E": 14})
    text(ws, "B2", "Session 26 · What is wrong with each chart?", TITLE)
    O = ["Title does not state the finding", "Bars do not start at zero", "Missing units", "Colours have no meaning", "No data source"]
    items = [("Q1", 'Line chart titled "Rainfall data" showing a clear drop in 2023.', "txt", O[0], O, None, None),
             ("Q2", "Bar chart whose y-axis runs from 800 to 1,000 mm.", "txt", O[1], O, None, None),
             ("Q3", 'y-axis labelled "Rain", 0 to 200, no unit or period.', "txt", O[2], O, None, None),
             ("Q4", "Anomaly chart with a random colour for each year.", "txt", O[3], O, None, None),
             ("Q5", "Clear chart of a district's rainfall with no hint of where the numbers came from.", "txt", O[4], O, None, None)]
    questions(ws, 4, items)
    ws = wb.create_sheet("S27 Chart choice")
    widths(ws, {"A": 3, "B": 5, "C": 66, "D": 32, "E": 14})
    text(ws, "B2", "Session 27 · Best chart for each question", TITLE)
    C = ["Line", "Bars by month", "Anomaly bars", "Bars by category", "Side-by-side bars", "Pie"]
    items = [("Q1", "How have rainy days per year changed since 1981?", "txt", "Line", C, None, None),
             ("Q2", "Which months bring most rain in an average year?", "txt", "Bars by month", C, None, None),
             ("Q3", "Was the last monsoon deficient compared to normal and to recent years?", "txt", "Anomaly bars", C, None, None),
             ("Q4", "How many days fall in each IMD intensity class?", "txt", "Bars by category", C, None, None),
             ("Q5", "NASA POWER vs ERA5 for each season of one year?", "txt", "Side-by-side bars", C, None, None),
             ("Q6", "Almost never best for comparing more than three values?", "txt", "Pie", C, None, None)]
    questions(ws, 4, items)
    # Dashboard
    ws = wb.create_sheet("S28 Dashboard")
    widths(ws, {"A": 3, "B": 24, "C": 18, "D": 24, "E": 18, "F": 24, "G": 18, "H": 3, "I": 8, "J": 12})
    text(ws, "B2", "Session 28 · Dashboard (updates with the Data tab)", TITLE)
    put(ws, "B3", "Headline (state a finding)", BOLD); ws.merge_cells("C3:G3"); put(ws, "C3", None, INP, YEL)
    kp = [("Normal year (mm)", "=Annual!K2", "#,##0"), ("Latest monsoon vs normal (%)", f"=INDEX({H_},COUNT({B_}))", "+0;-0"),
          ("Wettest day on record (mm)", f"=MAX({D_RANGE})", "0"), ("Rainy days per year", f"=AVERAGE({E_})", "0"),
          ("Wettest year", f"=INDEX(Annual!$A$2:$A${last},MATCH(MAX({B_}),{B_},0))", "0"), ("Driest year", f"=INDEX(Annual!$A$2:$A${last},MATCH(MIN({B_}),{B_},0))", "0")]
    for i, (k, f, fmt) in enumerate(kp):
        col = "BDF"[i % 3]; vcol = "CEG"[i % 3]; r = 5 + i // 3
        put(ws, f"{col}{r}", k, NOTE); c = put(ws, f"{vcol}{r}", f, Font(name=F, size=16, bold=True, color="2C6E9B"), GRN, fmt)
    hdr(ws, 8, ["Month", "Avg mm"], col=9)
    for m in range(1, 13):
        r = 8 + m; put(ws, f"I{r}", ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"][m - 1])
        put(ws, f"J{r}", f"=SUMIFS({D_RANGE},{M_RANGE},{m})/COUNT({B_})", fmt="0")
    c1 = BarChart(); c1.type = "col"; c1.title = "What is normal? Average rain by month (mm)"; c1.legend = None
    c1.add_data(Reference(ws, min_col=10, min_row=8, max_row=20), titles_from_data=True); c1.set_categories(Reference(ws, min_col=9, min_row=9, max_row=20))
    c1.height, c1.width = 7, 13; ws.add_chart(c1, "B9")
    c2 = BarChart(); c2.type = "col"; c2.title = "How were monsoons? % from normal"; c2.legend = None
    c2.add_data(Reference(at, min_col=8, min_row=1, max_row=last), titles_from_data=True); c2.set_categories(Reference(at, min_col=1, min_row=2, max_row=last))
    c2.height, c2.width = 7, 13; ws.add_chart(c2, "E9")
    c3 = LineChart(); c3.title = "Is anything changing? Rainy days per year"; c3.legend = None
    c3.add_data(Reference(at, min_col=5, min_row=1, max_row=last), titles_from_data=True); c3.set_categories(Reference(at, min_col=1, min_row=2, max_row=last))
    c3.height, c3.width = 7, 13; ws.add_chart(c3, "B24")
    put(ws, "E25", "Source note", BOLD); ws.merge_cells("E26:G28"); put(ws, "E26", "Data: (write your source). Normal = 1991–2020 average. Categories: IMD.", INP, YEL, wrap=True)
    # Story checker
    ws = wb.create_sheet("S29 Story checker")
    widths(ws, {"A": 3, "B": 44, "C": 80})
    text(ws, "B2", "Session 29 · Check my 150-word data story", TITLE)
    put(ws, "B4", "Paste your story here", BOLD); put(ws, "C4", None, INP, YEL, wrap=True); ws.row_dimensions[4].height = 160
    S = "C4"
    checks = [("Word count", f'=IF(LEN(TRIM({S}))=0,0,LEN(TRIM({S}))-LEN(SUBSTITUTE(TRIM({S})," ",""))+1)', None),
              ("120–180 words", f'=IF(AND(C6>=120,C6<=180),"✓ Correct","✗ Aim for 120–180")', True),
              ("Has a number", f'=IF(SUMPRODUCT(--ISNUMBER(FIND({{0,1,2,3,4,5,6,7,8,9}},{S})))>0,"✓ Correct","✗ Add a number")', True),
              ("Has a unit (mm, %, days)", f'=IF(OR(ISNUMBER(SEARCH("mm",{S})),ISNUMBER(SEARCH("%",{S})),ISNUMBER(SEARCH("days",{S}))),"✓ Correct","✗ Add a unit")', True),
              ("Makes a comparison", f'=IF(OR(ISNUMBER(SEARCH("normal",{S})),ISNUMBER(SEARCH("average",{S})),ISNUMBER(SEARCH("compared",{S})),ISNUMBER(SEARCH("than",{S}))),"✓ Correct","✗ Compare with normal, average or last year")', True),
              ("Names a source", f'=IF(OR(ISNUMBER(SEARCH("IMD",{S})),ISNUMBER(SEARCH("NASA",{S})),ISNUMBER(SEARCH("ERA5",{S})),ISNUMBER(SEARCH("CWC",{S})),ISNUMBER(SEARCH("data",{S}))),"✓ Correct","✗ Name your data source")', True)]
    for i, (k, f, chk) in enumerate(checks):
        r = 6 + i; put(ws, f"B{r}", k); put(ws, f"C{r}", f, BOLD)
    check_style(ws, "C7:C11")
    for s_ in ["Data", "Annual"]: wb.move_sheet(s_, offset=len(wb.sheetnames))
    finish(wb, "sheets/WDL_M6_practice.xlsx")


def finish(wb, path):
    import json, os
    reg = json.load(open("/tmp/wdl_registry.json")) if os.path.exists("/tmp/wdl_registry.json") else {}
    reg[path] = [r for r in REGISTRY if r[0] in wb.sheetnames]; REGISTRY.clear()
    json.dump(reg, open("/tmp/wdl_registry.json", "w"))
    for s in wb.worksheets:
        s.sheet_view.showGridLines = s.title in ("Data", "Annual", "Messy record", "Three sources 2021", "Reservoir data")
    wb.save(path); print("saved", path)


if __name__ == "__main__":
    import sys
    for m in sys.argv[1:] or ["2", "3", "4", "5", "6"]:
        globals()[f"build_m{m}"]()
