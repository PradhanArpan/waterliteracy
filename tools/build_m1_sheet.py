"""Module 1 practice workbook: the Sheet version of every practice activity, with self-check cells."""
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from nbkit import SAMPLE_GEN

exec(SAMPLE_GEN)
st = _make_sample_station()
N = len(st); LAST = N + 1

F = "Arial"
TITLE = Font(name=F, size=15, bold=True, color="13234A")
H = Font(name=F, size=10, bold=True, color="FFFFFF")
B = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="6B6B6B")
INP = Font(name=F, size=11, bold=True, color="0000FF")
HEAD = PatternFill("solid", fgColor="13234A")
YEL = PatternFill("solid", fgColor="FFF2CC")
GRN = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.border = H, HEAD, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def put(ws, ref, v, font=B, fill=None, fmt=None, wrap=False):
    c = ws[ref]; c.value = v; c.font = font; c.border = BOX
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if wrap: c.alignment = WRAP
    return c

def text(ws, ref, v, font=B):
    c = ws[ref]; c.value = v; c.font = font; c.alignment = WRAP

def widths(ws, w):
    for col, x in w.items(): ws.column_dimensions[col].width = x

def check_style(ws, rng):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0]},1)="✓"'], fill=GRN, font=Font(name=F, bold=True, color="2E7D32")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0]},1)="✗"'], fill=PatternFill("solid", fgColor="FDE2E1"), font=Font(name=F, bold=True, color="B71C1C")))

def num_check(cell, answer, tol="0.02"):
    return f'=IF({cell}="","",IF(ABS({cell}-({answer}))<=MAX(ABS({answer})*{tol},0.01),"✓ Correct","✗ Try again"))'

def txt_check(cell, answer):
    return f'=IF(OR({cell}="",{cell}="choose"),"",IF({cell}="{answer}","✓ Correct","✗ Try again"))'

def dv_list(ws, cells, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for c in cells: dv.add(c)

wb = Workbook()

# ---------------- Start here ----------------
ws = wb.active; ws.title = "Start here"
widths(ws, {"A": 3, "B": 100})
text(ws, "B2", "Module 1 · The Water Cycle as Data — Practice workbook", TITLE)
import os
UDEMY = os.environ.get("WDL_EDITION") == "udemy"
text(ws, "B3", "Water Data Literacy: Rainfall, Floods and Climate Change · Udemy edition · Dr Arpan Pradhan" if UDEMY else "Hydroinformatics in a Changing Climate: Water Data Literacy for All · CuX, CHRIST (Deemed to be University)", NOTE)
lines = [
    ("How to use this workbook", BOLD),
    ("This is the spreadsheet version of the Module 1 practice activities. It matches the Colab practice notebooks: same data, same questions, same answers.", B),
    ("1. Go to File → Make a copy so you have your own version.", B),
    ("2. Each tab is one session: S01 to S04 are practice activities, S05 is the workshop.", B),
    ("3. Type or choose your answers in the YELLOW cells. The CHECK column tells you if you are right (✓) or should try again (✗).", B),
    ("4. The Data tab holds the sample station record. Charts and totals are calculated from it with formulas.", B),
    ("", B),
    ("About the data", BOLD),
    ("The sample station is ILLUSTRATIVE, SYNTHETIC data for practice, not an official IMD or CWC record. In Module 2 you will download real data for your own district.", B),
]
for i, (t, f) in enumerate(lines):
    text(ws, f"B{5+i}", t, f)

# ---------------- Data ----------------
dt = wb.create_sheet("Data")
widths(dt, {"A": 13, "B": 11, "C": 13, "D": 8})
hdr(dt, 1, ["Date", "Rain (mm)", "River level (m)", "Month"])
for i, row in st.iterrows():
    r = i + 2
    put(dt, f"A{r}", row.date.to_pydatetime(), fmt="dd-mmm-yyyy")
    put(dt, f"B{r}", float(row.rain_mm), fmt="0.0")
    put(dt, f"C{r}", float(row.level_m), fmt="0.00")
    put(dt, f"D{r}", f"=MONTH(A{r})")
dt.freeze_panes = "A2"
RA, LA, DA, MA = f"Data!$B$2:$B${LAST}", f"Data!$C$2:$C${LAST}", f"Data!$A$2:$A${LAST}", f"Data!$D$2:$D${LAST}"

# ---------------- S01 ----------------
ws = wb.create_sheet("S01 Headlines")
widths(ws, {"A": 3, "B": 6, "C": 56, "D": 24, "E": 16})
text(ws, "B2", "Session 1 · What was measured?", TITLE)
text(ws, "B3", "Choose the kind of water number each example headline reports.", NOTE)
hdr(ws, 5, ["#", "Headline (example)", "Your answer", "Check"], col=2)
opts = ["Rainfall depth", "Water level", "Water supply", "Rainfall anomaly", "River flow", "Groundwater depth"]
heads = [("A", "Heavy rain: 87 mm recorded at the city observatory", "Rainfall depth"),
         ("B", "Lake level drops 1.2 m in three months", "Water level"),
         ("C", "Water board to cut supply from 900 MLD to 750 MLD", "Water supply"),
         ("D", "Monsoon 9% above normal so far", "Rainfall anomaly"),
         ("E", "Dam outflow raised to 30,000 cusecs", "River flow"),
         ("F", "Groundwater at 18 m below ground in summer", "Groundwater depth")]
for i, (k, h, a) in enumerate(heads):
    r = 6 + i
    put(ws, f"B{r}", k, BOLD); put(ws, f"C{r}", h, wrap=True)
    put(ws, f"D{r}", None, INP, YEL); put(ws, f"E{r}", txt_check(f"D{r}", a), BOLD)
dv_list(ws, [f"D{6+i}" for i in range(6)], opts)
put(ws, "C12", "Score (out of 6)", BOLD); put(ws, "E12", '=COUNTIF(E6:E11,"✓*")', BOLD, GRN)
text(ws, "B14", "Task 2", BOLD)
put(ws, "C15", 'Headline: "Reservoir storage at 42% of capacity". Which question does it answer least well?', wrap=True)
put(ws, "D15", None, INP, YEL); put(ws, "E15", txt_check("D15", "Compared to what?"), BOLD)
dv_list(ws, ["D15"], ["What was measured?", "Where?", "Over what time?", "Compared to what?"])
ws.row_dimensions[15].height = 30
check_style(ws, "E6:E15")

# ---------------- S02 ----------------
ws = wb.create_sheet("S02 Volumes")
widths(ws, {"A": 3, "B": 6, "C": 62, "D": 16, "E": 16, "F": 46})
text(ws, "B2", "Session 2 · Rain is a volume; water units", TITLE)
text(ws, "B3", "Work each one out, type your answer, and check. Answers within 2% count as correct.", NOTE)
hdr(ws, 5, ["#", "Question", "Your answer", "Check", "Hint"], col=2)
qs = [
    ("A 120 m² roof receives 65 mm of rain. How many litres fall on it?", "65*120", "Depth (mm) × area (m²) = litres."),
    ("A 30 ha campus receives 40 mm of rain. How many m³ fall on it?", "40*30*10000/1000", "1 ha = 10,000 m². Litres ÷ 1,000 = m³."),
    ("A reservoir releases 2 TMC. How many MCM is that?", "2*28.316846592", "1 TMC ≈ 28.3 MCM."),
    ("A river flows at 250 m³/s for one day. How many MCM pass?", "250*86400/1000000", "Flow × 86,400 seconds = m³; ÷ 1,000,000 = MCM."),
    ("A dam releases 50,000 cusecs. What is this in m³/s?", "50000/35.3147", "1 m³/s ≈ 35.3 cusecs."),
    ("A city uses 1,450 MLD. For how many days would 1 TMC last?", "28316.846592/1450", "1 TMC ≈ 28,317 million litres."),
]
for i, (q, ans, hint) in enumerate(qs):
    r = 6 + i
    put(ws, f"B{r}", f"Q{i+1}", BOLD); put(ws, f"C{r}", q, wrap=True)
    put(ws, f"D{r}", None, INP, YEL, "#,##0.00"); put(ws, f"E{r}", num_check(f"D{r}", ans), BOLD)
    put(ws, f"F{r}", hint, NOTE, wrap=True); ws.row_dimensions[r].height = 30
put(ws, "C12", "Score (out of 6)", BOLD); put(ws, "E12", '=COUNTIF(E6:E11,"✓*")', BOLD, GRN)
check_style(ws, "E6:E11")
text(ws, "B14", "Calculator: rain to volume", BOLD)
hdr(ws, 15, ["", "Input", "Value"], col=2)
put(ws, "C16", "Rainfall (mm)"); put(ws, "D16", 50, INP, YEL)
put(ws, "C17", "Area (m²)"); put(ws, "D17", 100, INP, YEL, "#,##0")
put(ws, "C18", "Volume (litres)", BOLD); put(ws, "D18", "=D16*D17", BOLD, GRN, "#,##0")
put(ws, "C19", "Volume (MCM)", BOLD); put(ws, "D19", "=D18/1E9", BOLD, GRN, "0.000000")

# ---------------- S03 ----------------
ws = wb.create_sheet("S03 Rain Gauge")
widths(ws, {"A": 3, "B": 6, "C": 58, "D": 22, "E": 16, "F": 3, "G": 14, "H": 13})
text(ws, "B2", "Session 3 · Reading a rain gauge record", TITLE)
text(ws, "B3", "Use the monthly table and the Data tab to answer. Rainy day = 2.5 mm or more (IMD).", NOTE)
hdr(ws, 5, ["#", "Question", "Your answer", "Check"], col=2)
q3 = [
    ("Total rainfall in 2023 (mm)", num_check("D6", f"SUM({RA})"), None),
    ("Number of rainy days (2.5 mm or more)", num_check("D7", f'COUNTIF({RA},">=2.5")', "0"), None),
    ("Wettest month", None, "Sep"),
    ("% of the year's rain on the 10 wettest days", num_check("D9", f"SUMPRODUCT(LARGE({RA},ROW($1:$10)))/SUM({RA})*100", "0.05"), None),
    ("Season with the most rain", None, "SW monsoon (Jun–Sep)"),
]
for i, (q, f, a) in enumerate(q3):
    r = 6 + i
    put(ws, f"B{r}", f"Q{i+1}", BOLD); put(ws, f"C{r}", q, wrap=True)
    put(ws, f"D{r}", None, INP, YEL); put(ws, f"E{r}", f if f else txt_check(f"D{r}", a), BOLD)
months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
dv_list(ws, ["D8"], months)
dv_list(ws, ["D10"], ["Winter (Jan–Feb)", "Pre-monsoon (Mar–May)", "SW monsoon (Jun–Sep)", "NE monsoon (Oct–Dec)"])
check_style(ws, "E6:E10")
hdr(ws, 5, ["Month", "Rain (mm)"], col=7)
for i, m in enumerate(months):
    r = 6 + i
    put(ws, f"G{r}", m); put(ws, f"H{r}", f"=SUMIF({MA},{i+1},{RA})", BOLD, GRN, "#,##0.0")
put(ws, "G18", "Year", BOLD); put(ws, "H18", "=SUM(H6:H17)", BOLD, GRN, "#,##0.0")
ch = BarChart(); ch.type = "col"; ch.title = "Monthly rainfall, 2023 (sample station)"; ch.legend = None
ch.y_axis.title = "Rainfall (mm)"
ch.add_data(Reference(ws, min_col=8, min_row=5, max_row=17), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=7, min_row=6, max_row=17)); ch.height = 7.5; ch.width = 16
ws.add_chart(ch, "B13")

# ---------------- S04 ----------------
ws = wb.create_sheet("S04 River Gauge")
widths(ws, {"A": 3, "B": 6, "C": 62, "D": 16, "E": 16, "F": 3, "G": 14, "H": 11, "I": 13})
text(ws, "B2", "Session 4 · Reading a river gauge record", TITLE)
text(ws, "B3", "Warning level 3.6 m, danger level 4.2 m. Use the flood table on the right.", NOTE)
put(ws, "C4", "Warning level (m)"); put(ws, "D4", 3.6, BOLD)
hdr(ws, 5, ["#", "Question", "Your answer", "Check"], col=2)
pk = f"INDEX({DA},MATCH(MAX({LA}),{LA},0))"
wet = f"INDEX({DA},MATCH(MAX({RA}),{RA},0))"
q4 = [
    ("Highest river level of the year (m)", num_check("D6", f"MAX({LA})", "0.005")),
    ("Lag time: days between the heaviest rain and the river peak", num_check("D7", f"{pk}-{wet}", "0")),
    ("Days above the danger level (4.2 m)", num_check("D8", f'COUNTIF({LA},">4.2")', "0")),
    ("Days from the peak until the river fell back below the warning level", num_check("D9", "5", "0")),
]
for i, (q, f) in enumerate(q4):
    r = 6 + i
    put(ws, f"B{r}", f"Q{i+1}", BOLD); put(ws, f"C{r}", q, wrap=True)
    put(ws, f"D{r}", None, INP, YEL); put(ws, f"E{r}", f, BOLD)
check_style(ws, "E6:E9")
hdr(ws, 5, ["Date", "Rain (mm)", "Level (m)"], col=7)
start = st.index[st.date == pd.Timestamp("2023-09-15")][0] + 2
for k in range(26):
    r, src = 6 + k, start + k
    put(ws, f"G{r}", f"=Data!A{src}", fmt="dd-mmm")
    put(ws, f"H{r}", f"=Data!B{src}", fmt="0.0")
    put(ws, f"I{r}", f"=Data!C{src}", fmt="0.00")
ws.conditional_formatting.add("I6:I31", FormulaRule(formula=["I6>4.2"], fill=PatternFill("solid", fgColor="F4C7C3")))
ws.conditional_formatting.add("I6:I31", FormulaRule(formula=["I6>3.6"], fill=PatternFill("solid", fgColor="FCE4D6")))
lc = LineChart(); lc.title = "River level, 15 Sep to 10 Oct"; lc.legend = None; lc.y_axis.title = "Level (m)"
lc.add_data(Reference(ws, min_col=9, min_row=5, max_row=31), titles_from_data=True)
lc.set_categories(Reference(ws, min_col=7, min_row=6, max_row=31)); lc.height = 7.5; lc.width = 16
lc.x_axis.number_format = "dd-mmm"
ws.add_chart(lc, "B12")

# ---------------- S05 ----------------
ws = wb.create_sheet("S05 District Card")
widths(ws, {"A": 3, "B": 46, "C": 26, "D": 48})
text(ws, "B2", "Session 5 · Workshop: My district in numbers", TITLE)
text(ws, "B3", "Replace the example values with your own district's facts. This card is step 1 of your capstone.", NOTE)
hdr(ws, 5, ["My district", "Value", "Where to find it"], col=2)
facts = [("District", "Example district", "Your choice", None), ("State", "Example state", "", None),
         ("Area (km²)", 2000, "District website or Census of India", "#,##0"),
         ("Population", 2000000, "Census of India or district website", "#,##0"),
         ("Normal annual rainfall (mm)", 900, "IMD district rainfall pages", "#,##0"),
         ("Main river or lake", "Example river", "Map or district website", None)]
for i, (lab, v, src, fmt) in enumerate(facts):
    r = 6 + i
    put(ws, f"B{r}", lab); put(ws, f"C{r}", v, INP, YEL, fmt); put(ws, f"D{r}", src, NOTE)
hdr(ws, 13, ["In a normal year", "Value", "How it is calculated"], col=2)
calc = [("Rain on the district (MCM)", "=C10/1000*C8*1000000/1000000", "#,##0", "rainfall (m) × area (m²) ÷ 1,000,000"),
        ("Rain on the district (TMC)", "=C14/28.316846592", "#,##0.0", "MCM ÷ 28.3"),
        ("Litres per person per day, if all rain were shared", "=C14*1E9/C9/365", "#,##0", "litres ÷ people ÷ 365"),
        ("Domestic need at 135 LPCD (MCM per year)", "=C9*135*365/1E9", "#,##0", "people × 135 L × 365 days"),
        ("Rain received ÷ water needed", "=C14/C17", "0.0", "how many times the need")]
for i, (lab, f, fmt, how) in enumerate(calc):
    r = 14 + i
    put(ws, f"B{r}", lab); put(ws, f"C{r}", f, BOLD, GRN, fmt); put(ws, f"D{r}", how, NOTE)
hdr(ws, 20, ["Lab 1 question", "Your answer", ""], col=2)
labq = ["Why do cities flood quickly?", "Most useful time scale for a farmer, and why", "Why does the river peak after the rain?"]
for i, q in enumerate(labq):
    r = 21 + i
    put(ws, f"B{r}", q, wrap=True); ws.merge_cells(f"C{r}:D{r}")
    put(ws, f"C{r}", None, INP, YEL, wrap=True); ws.row_dimensions[r].height = 42
put(ws, "B25", "Card status", BOLD)
put(ws, "C25", '=IF(C6="Example district","✗ Replace the example values with your own district",IF(COUNTA(C21:C23)<3,"✗ Answer all three Lab 1 questions","✓ Ready: take a screenshot of this tab for Lab 1"))', BOLD)
ws.merge_cells("C25:D25")
check_style(ws, "C25:C25")
bc = BarChart(); bc.type = "bar"; bc.title = "Rain received vs domestic need (MCM per year)"; bc.legend = None
put(ws, "F14", "Need at 135 LPCD"); put(ws, "G14", "=C17", fmt="#,##0")
put(ws, "F15", "Rain on the district"); put(ws, "G15", "=C14", fmt="#,##0")
bc.add_data(Reference(ws, min_col=7, min_row=14, max_row=15)); bc.set_categories(Reference(ws, min_col=6, min_row=14, max_row=15))
bc.height = 6; bc.width = 16
ws.add_chart(bc, "B27")

for s in wb.worksheets:
    s.sheet_view.showGridLines = s.title == "Data"
wb.move_sheet("Data", offset=len(wb.sheetnames))
os.makedirs("udemy/sheets" if UDEMY else "sheets", exist_ok=True)
wb.save(("udemy/sheets" if UDEMY else "sheets") + "/WDL_M1_practice.xlsx")
print("saved")
