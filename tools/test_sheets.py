"""Fill every answer cell with its correct answer, recalculate with LibreOffice, and count ✓ / ✗ / errors."""
import json, subprocess, shutil, sys
from openpyxl import load_workbook
reg = json.load(open("/tmp/wdl_registry.json"))
for path, items in reg.items():
    tmp = "/tmp/test_" + path.split("/")[-1]
    wb = load_workbook(path)
    for sheet, cell, kind, ans in items:
        wb[sheet][cell] = ("=" + ans) if kind == "num" else (ans if not ans.startswith("=") else ans)
    wb.save(tmp)
    out = subprocess.run([sys.executable, "/mnt/skills/public/xlsx/scripts/recalc.py", tmp, "120"], capture_output=True, text=True).stdout
    res = json.loads(out)
    wv = load_workbook(tmp, data_only=True)
    checks = []
    for sheet, cell, kind, ans in items:
        col = chr(ord(cell[0]) + 1); v = wv[sheet][f"{col}{cell[1:]}"].value
        checks.append((sheet, cell, v))
    good = sum(str(v).startswith("✓") for _, _, v in checks)
    print(f"{path}: {good}/{len(checks)} ✓ · formula errors {res.get('total_errors')} {res.get('error_summary') or ''}")
    for s, c, v in checks:
        if not str(v).startswith("✓"): print("    ", s, c, v)
