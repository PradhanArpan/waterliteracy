"""Udemy curriculum workbook: every section and item in upload order, with lecture resources and links."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from course import MODULES, sessions, DUR

GH = "https://colab.research.google.com/github/PradhanArpan/waterliteracy/blob/main/udemy/notebooks"
RAW = "https://github.com/PradhanArpan/waterliteracy/raw/main/udemy/sheets"
src = open("tools/make_sessions_json.py").read()
REFS = eval(src[src.index("REFS = {") + 7: src.index("}\n", src.index("REFS = {")) + 1])

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=10); HEAD = PatternFill("solid", fgColor="13234A")
SEC = PatternFill("solid", fgColor="DCE6F1"); B = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
LINK = Font(name=F, size=10, color="1155CC", underline="single")

rows = []   # (section, item no, type, title, minutes, preview, resources[list of (label, url)], notes)
S = sessions()
item = 0
for m in MODULES:
    sec = f"Module {m['n']}: {m['title']}"
    item += 1; rows.append((sec, item, "Video", f"Introduction to {m['title']}", DUR["intro"], m["n"] == 1, [], "Face to camera"))
    for s in [x for x in S if x["module"] == m["n"]]:
        k = s["num"]; base = f"{GH}/M{m['n']}/WDL_M{m['n']}_S{k:02d}"
        sheet = (f"Practice workbook (Excel / Google Sheets), Module {m['n']}", f"{RAW}/WDL_M{m['n']}_practice.xlsx")
        if not s["workshop"]:
            res = [("Reading notebook (Colab)", f"{base}_reading.ipynb"), ("Practice notebook (Colab)", f"{base}_practice.ipynb"), sheet]
            if k in REFS: res.append(("Reference: " + REFS[k][0], REFS[k][1]))
            prev = k in (1, 2)
            item += 1; rows.append((sec, item, "Video", f"{k}.1 {s['title']}: {s['hook']}", DUR["concept"], prev, [], "Concept · face to camera"))
            item += 1; rows.append((sec, item, "Video", f"{k}.2 {s['title']}: notebook walkthrough", DUR["walkthrough"], prev and k == 1, res,
                                    "Screen recording. Attach the resources listed."))
        else:
            res = [("Workshop notebook (Colab)", f"{base}_workshop.ipynb"), sheet]
            item += 1; rows.append((sec, item, "Video", f"{k}.1 Workshop briefing: {s['title'].replace('Workshop: ', '').replace('Capstone workshop and submission', 'My District’s Water Story')}",
                                    DUR["ws_video"], False, res, "Screen recording"))
            lab = "Capstone: My District's Water Story" if m["n"] == 6 else f"Lab {m['n']}: {s['title'].replace('Workshop: ', '')}"
            item += 1; rows.append((sec, item, "Assignment", lab, 40 if m["n"] < 6 else 60, False, [], "See the Assignments tab of the course kit"))
            item += 1; rows.append((sec, item, "Quiz", f"Module {m['n']} quiz", 15, False, [], f"10 questions: scripts doc, tab 'Module {m['n']} Quiz'"))
rows.append(("Module 6: From Data to Dashboard", item + 1, "Video", "Thank you and what next", 2, False, [], "Face to camera"))

wb = Workbook(); ws = wb.active; ws.title = "Curriculum"
cols = ["Section", "#", "Type", "Title", "Minutes", "Free preview", "Resources (add as external links)", "Notes"]
widths = [34, 5, 11, 58, 9, 10, 70, 32]
for i, (c, w) in enumerate(zip(cols, widths), 1):
    cell = ws.cell(1, i, c); cell.font, cell.fill, cell.border = H, HEAD, BOX
    ws.column_dimensions[chr(64 + i)].width = w
r = 2; prev_sec = None
for sec, n, typ, title, mins, prev, res, note in rows:
    if sec != prev_sec:
        c = ws.cell(r, 1, sec); c.font = BOLD
        for j in range(1, 9): ws.cell(r, j).fill = SEC; ws.cell(r, j).border = BOX
        r += 1; prev_sec = sec
    vals = ["", n, typ, title, mins, "Yes" if prev else "", "\n".join(f"{l}: {u}" for l, u in res), note]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = B; c.border = BOX; c.alignment = WRAP
    if res: ws.row_dimensions[r].height = 15 * len(res) + 4
    r += 1
tot_v = sum(x[4] for x in rows if x[2] == "Video"); tot = sum(x[4] for x in rows)
ws.cell(r + 1, 4, "Video minutes").font = BOLD; ws.cell(r + 1, 5, tot_v).font = BOLD
ws.cell(r + 2, 4, "All items, minutes (incl. assignments and quizzes)").font = BOLD; ws.cell(r + 2, 5, tot).font = BOLD
ws.cell(r + 3, 4, "Free preview minutes").font = BOLD; ws.cell(r + 3, 5, sum(x[4] for x in rows if x[5])).font = BOLD
ws.freeze_panes = "A2"

# flat link list: one row per resource (easy to copy into Udemy's "External resource" dialog)
ls = wb.create_sheet("Resource links")
for i, (c, w) in enumerate(zip(["Lecture #", "Lecture title", "Resource title", "URL"], [10, 58, 52, 110]), 1):
    cell = ls.cell(1, i, c); cell.font, cell.fill = H, HEAD; ls.column_dimensions[chr(64 + i)].width = w
k = 2
for sec, n, typ, title, mins, prev, res, note in rows:
    for l, u in res:
        ls.cell(k, 1, n).font = B; ls.cell(k, 2, title).font = B; ls.cell(k, 3, l).font = B
        c = ls.cell(k, 4, u); c.hyperlink = u; c.font = LINK; k += 1
ls.freeze_panes = "A2"
wb.save("udemy/Udemy_curriculum.xlsx")
nv = sum(1 for x in rows if x[2] == "Video")
print(f"items {len(rows)} · videos {nv} ({tot_v} min) · assignments {sum(1 for x in rows if x[2]=='Assignment')} · quizzes {sum(1 for x in rows if x[2]=='Quiz')} · preview {sum(x[4] for x in rows if x[5])} min · links {k-2}")
